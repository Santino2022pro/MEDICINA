# -*- coding: utf-8 -*-
"""Componentes visuales con la identidad de la serie (etiquetas amarillas,
recuadros con borde oscuro, insignias, moléculas, enzimas, tablas)."""
from lib import *
import math


def lighten(c, f):
    r, g, b = int(c[0:2], 16), int(c[2:4], 16), int(c[4:6], 16)
    return '%02X%02X%02X' % tuple(int(v + (255 - v) * f) for v in (r, g, b))


def darken(c, f):
    r, g, b = int(c[0:2], 16), int(c[2:4], 16), int(c[4:6], 16)
    return '%02X%02X%02X' % tuple(int(v * (1 - f)) for v in (r, g, b))


def shade(c, hl=0.5, dk=0.28, focus=(32, 28)):
    return ('rad', [(0, lighten(c, hl)), (55, c), (100, darken(c, dk))], focus)


def vshade(c, hl=0.35, dk=0.3):
    """sombreado cilíndrico (lineal) para fibras y tubos"""
    return ('lin', [(0, darken(c, dk)), (30, lighten(c, hl)), (55, c), (100, darken(c, dk))], 90)


SHD = (5, 3, 32)

# ---------------------------------------------------------------- encabezados
UNIT_LABEL = {}


PROTECT = {'Ca', 'Gs', 'Gi', 'Gq', 'Akt', 'Na', 'Mg', 'pH', 'Km', 'Vm', 'mV', 'Epac2', 'n', 'mOsm', 'kg', 'g', 'h', 'mg', 'dL', 'L', 'Lp', 'mTORC1'}
GREEK = 'αβγδεκλμπστφωΔ'


def smart_upper(text):
    """mayúsculas que preservan símbolos científicos (AMPc, Ca, Gs, CoA, letras griegas, sub/superíndices)"""
    out = []
    for tok in re.split(r'(\^\{[^}]*\}|_\{[^}]*\}|[A-Za-zÁÉÍÓÚÜÑáéíóúüñ0-9αβγδεκλμπστφωΔ]+)', text):
        if not tok or tok.startswith('^{') or tok.startswith('_{'):
            out.append(tok); continue
        if re.fullmatch(r'[A-Za-zÁÉÍÓÚÜÑáéíóúüñ0-9αβγδεκλμπστφωΔ]+', tok):
            if tok in PROTECT or re.search(r'[a-zñ].*[A-Z]|[A-Z]{2,}.*[a-z]|^[a-z]+[A-Z0-9]|^\d+[a-z]+$', tok):
                out.append(tok); continue
            out.append(''.join(c if c in GREEK else c.upper() for c in tok)); continue
        out.append(tok)
    return ''.join(out)


def title(s, text, tag=None, size=26):
    """etiqueta de título (izquierda) y etiqueta de sección (derecha)"""
    t = smart_upper(text)
    tagw = 0
    if tag:
        tsz = 13
        tagw = text_w(plain(smart_upper(tag)), F_HEAD, True, tsz) + 0.4
        tagw = min(tagw, 3.6)
        s.shape(SW - 0.3 - tagw, 0.22, tagw, 0.46, 'rect', TAG, (INK, 1.5), SHD, smart_upper(tag), tsz, F_HEAD, INK, True,
                'c', 'm', (0.08, 0.02, 0.08, 0.02), name='Etiqueta sección', fit=True, minsize=10)
    maxw = SW - 0.6 - tagw - 0.45
    sz = size
    while text_w(plain(t), F_HEAD, True, sz) + 0.45 > maxw and sz > 17:
        sz -= 0.5
    w = min(text_w(plain(t), F_HEAD, True, sz) + 0.45, maxw)
    lines = 1 if text_w(plain(t), F_HEAD, True, sz) + 0.4 <= maxw else 2
    h = 0.62 if lines == 1 else 0.95
    return s.shape(0.3, 0.16, w, h, 'rect', TAG, (INK, 1.5), SHD, t, sz, F_HEAD, INK, True, 'c', 'm',
                   (0.12, 0.02, 0.12, 0.02), name='Título', fit=True, minsize=15)


def pagenum(s, n):
    s.shape(SW - 0.62, SH - 0.42, 0.42, 0.28, 'roundRect', ('a', 'FFFFFF', 55), None, False, str(n), 10, F_HEAD,
            '5A3A2A', True, 'c', 'm', (0, 0, 0, 0), name='Número', adj={'adj': 50000})


# ---------------------------------------------------------------- recuadros
def card(s, x, y, w, h, text, fill=SALMON, size=15, color=INK, line=(INK, 1.25), shadow=SHD, geom='rect',
         anchor='t', align='l', name=None, minsize=11.5, head=None, headfill=TAG, headsize=None, spc=4, bold=False, lnspc=None,
         grow=True):
    """recuadro de información. Si head: chip de encabezado superpuesto arriba a la izquierda (agrupado)."""
    hs0 = headsize or min(size + 1, 16)
    if grow:
        size = min(size + 2.5, max(size, 18))
    if head:
        g = s.group_start(name or ('Recuadro ' + plain(head)[:25]))
        s.shape(x, y + 0.2, w, h - 0.2, geom, fill, line, shadow, text, size, F_BODY, color, bold, align, anchor,
                (0.14, 0.3, 0.14, 0.08), minsize=minsize, spc=spc, lnspc=lnspc)
        hs = hs0
        hw = min(text_w(plain(smart_upper(head)), F_HEAD, True, hs) + 0.36, w - 0.2)
        s.shape(x + 0.12, y, hw, 0.38, 'rect', headfill, (INK, 1.25), False, smart_upper(head), hs, F_HEAD,
                WHITE if headfill in (RED, INH, TRN, ACT, '7030A0', NIT, '2E4053') else INK, True, 'c', 'm', (0.05, 0, 0.05, 0), minsize=10)
        return s.group_end()
    return s.shape(x, y, w, h, geom, fill, line, shadow, text, size, F_BODY, color, bold, align, anchor,
                   (0.14, 0.08, 0.14, 0.08), name=name, minsize=minsize, spc=spc, lnspc=lnspc)


def note_box(s, x, y, w, h, text, size=14, name='Perla'):
    """'perla' / advertencia conceptual: rojo con texto blanco (como en la serie)"""
    return card(s, x, y, w, h, text, RED, size, WHITE, (INK, 1.5), SHD, name=name)


def chip(s, x, y, text, fill=TAG, size=14, color=INK, h=None, name=None, w=None, line=(INK, 1.25), font=F_HEAD,
         bold=True, shadow=False, geom='rect', align='c'):
    tw = text_w(plain(text), font, bold, size) + 0.26
    w = w or tw
    h = h or size / 72 * 1.55
    return s.shape(x, y, w, h, geom, fill, line, shadow, text, size, font, color, bold, align, 'm', (0.05, 0, 0.05, 0),
                   name=name, minsize=9)


def chipw(text, size=14, font=F_HEAD, bold=True):
    return text_w(plain(text), font, bold, size) + 0.26


def label(s, x, y, w, h, text, size=13, color=INK, bold=True, align='c', font=F_HEAD, name=None, anchor='m'):
    return s.text(x, y, w, h, text, size, font, color, bold, align, anchor, name=name, minsize=9)


def badge(s, x, y, n, d=0.42, fill=TAG, color=INK, name=None):
    return s.shape(x, y, d, d, 'ellipse', fill, (INK, 1.5), SHD, str(n), 15 if d >= 0.4 else 12, F_HEAD, color, True, 'c', 'm',
                   (0, 0, 0, 0), name=name or ('Paso %s' % n), fit=False)


def enz(s, x, y, text, w=None, fill='FFFFFF', size=12.5, color='1F3864', name=None, h=0.32):
    """enzima: píldora blanca con texto azul oscuro en negrita"""
    w = w or (text_w(plain(text), F_HEAD, True, size) + 0.36)
    return s.shape(x, y, w, h, 'roundRect', fill, ('1F3864', 1.25), False, text, size, F_HEAD, color, True, 'c', 'm',
                   (0.03, 0, 0.03, 0), name=name or ('Enzima ' + plain(text)[:20]), adj={'adj': 50000}, minsize=9)


def sign(s, x, y, kind='+', d=0.3, name=None):
    col = ACT if kind == '+' else INH
    return s.shape(x, y, d, d, 'ellipse', col, (WHITE, 1.25), False, '+' if kind == '+' else '−', 14, F_HEAD, WHITE, True,
                   'c', 'm', (0, 0, 0, 0), name=name or ('Signo ' + kind), fit=False)


def mol(s, x, y, key, text=None, d=0.42, name=None, size=None, w=None):
    """ficha molecular coherente por color (glucosa hexágono, AG zigzag, resto círculo/píldora)"""
    fillc, tc = MOL[key]
    geom = 'hexagon' if key in ('glc', 'fru') else ('ellipse' if not w else 'roundRect')
    text = text if text is not None else {'glc': 'Glc', 'lac': 'Lac', 'ala': 'Ala', 'ag': 'AG', 'gly': 'Gli', 'cc': 'CC',
                                          'nh4': 'NH_{4}^{+}', 'urea': 'Urea', 'gln': 'Gln', 'aa': 'AA', 'ins': 'Ins',
                                          'gcg': 'Gcg', 'atp': 'ATP', 'ca': 'Ca^{2+}', 'k': 'K^{+}', 'tag': 'TAG',
                                          'chol': 'Col', 'o2': 'O_{2}', 'co2': 'CO_{2}', 'urate': 'Urato', 'fru': 'Fru',
                                          'acoa': 'Ac-CoA', 'mal': 'Mal-CoA'}[key]
    ww = w or d
    size = size or max(8.5, min(13, d * 26 - (len(plain(text)) - 3) * 1.2))
    return s.shape(x, y, ww, d, geom, shade(fillc, 0.35, 0.2), (darken(fillc, 0.45), 1.0), False, text, size, F_HEAD, tc, True,
                   'c', 'm', (0, 0, 0, 0), name=name or ('Molécula ' + key), fit=False,
                   adj={'adj': 50000} if geom == 'roundRect' else ({'adj': 28000, 'vf': 115470} if geom == 'hexagon' else None))


def ag_chain(s, x, y, w=0.9, h=0.22, name=None, color='B7950B'):
    """ácido graso: zigzag con grupo carboxilo"""
    g = s.group_start(name or 'Ácido graso')
    n = 8
    pts = []
    for i in range(n + 1):
        pts.append((x + 0.16 + (w - 0.16) * i / n, y + (h if i % 2 else 0.02)))
    s.arrow(pts, color, 2.25, tail=None)
    s.shape(x - 0.02, y - 0.04, 0.2, 0.2, 'ellipse', shade('E74C3C', 0.4, 0.2), (darken('E74C3C', 0.4), 0.75), False)
    return s.group_end()


def arrow_lbl(s, pts, text=None, color=INK, w=2.25, size=12, lx=None, ly=None, lw=1.6, curve=False, dash=None, name=None,
              lcolor=None, tail='triangle', head=None, lbg=None):
    """flecha con etiqueta opcional (devuelve grupo si hay etiqueta)"""
    if not text:
        return s.arrow(pts, color, w, head=head, tail=tail, curve=curve, dash=dash, name=name)
    g = s.group_start(name or ('Flecha ' + plain(text)[:18]))
    s.arrow(pts, color, w, head=head, tail=tail, curve=curve, dash=dash)
    if lx is None:
        mx = sum(p[0] for p in pts) / len(pts); my = sum(p[1] for p in pts) / len(pts)
        lx, ly = mx - lw / 2, my - 0.32
    nl = wrap_lines(text, lw, size, F_HEAD, True)
    hh = nl * size * 1.2 / 72 + 0.04
    if lbg:
        s.shape(lx, ly, lw, hh, 'roundRect', lbg, None, False, text, size, F_HEAD, lcolor or color, True, 'c', 'm',
                (0.02, 0, 0.02, 0), minsize=9, adj={'adj': 30000})
    else:
        s.text(lx, ly, lw, hh, text, size, F_HEAD, lcolor or color, True, 'c', 'm', minsize=9)
    return s.group_end()


def flow(s, items, x, y, w, h, fills=None, size=14, gap=0.35, dir='h', acolor=INK, names=None, minsize=11, geom='rect',
         shadow=SHD, line=(INK, 1.25), bold=False):
    """cadena de recuadros conectados por flechas; devuelve [(id_caja, id_flecha_entrante)]"""
    n = len(items)
    res = []
    if dir == 'h':
        bw = (w - gap * (n - 1)) / n
        for i, t in enumerate(items):
            bx = x + i * (bw + gap)
            f = (fills[i] if isinstance(fills, list) else fills) or SALMON
            a = None
            if i:
                a = s.arrow([(bx - gap + 0.04, y + h / 2), (bx - 0.04, y + h / 2)], acolor, 2.5)
            b = s.shape(bx, y, bw, h, geom, f, line, shadow, t, size, F_BODY, WHITE if f in (RED, INH, TRN, ACT) else INK, bold,
                        'c', 'm', (0.06, 0.04, 0.06, 0.04), minsize=minsize, name=(names[i] if names else None),
                        adj={'adj': 12000} if geom == 'roundRect' else None)
            res.append((b, a))
    else:
        bh = (h - gap * (n - 1)) / n
        for i, t in enumerate(items):
            by = y + i * (bh + gap)
            f = (fills[i] if isinstance(fills, list) else fills) or SALMON
            a = None
            if i:
                a = s.arrow([(x + w / 2, by - gap + 0.03), (x + w / 2, by - 0.03)], acolor, 2.5)
            b = s.shape(x, by, w, bh, geom, f, line, shadow, t, size, F_BODY, WHITE if f in (RED, INH, TRN, ACT) else INK, bold,
                        'c', 'm', (0.08, 0.03, 0.08, 0.03), minsize=minsize, name=(names[i] if names else None),
                        adj={'adj': 12000} if geom == 'roundRect' else None)
            res.append((b, a))
    return res


def flow_anim(s, res, first_click=True):
    """anima una cadena: cada caja con su flecha en un clic"""
    for b, a in res:
        if a:
            s.step(wipe(a, 'l'), fade(b, after=True))
        else:
            s.step(fade(b))


# ---------------------------------------------------------------- tablas nativas
def table(s, x, y, w, rows, colw=None, rowh=0.4, size=12.5, head_fill=TAG, fills=(CREAM, LSALMON), name=None,
          first_col_fill=None, head_size=None, head_color=INK, bold_first=True, heights=None):
    """tabla editable de PowerPoint; las celdas aceptan el marcado de texto enriquecido"""
    ncol = len(rows[0]); nrow = len(rows)
    colw = colw or [w / ncol] * ncol
    tot = sum(colw); colw = [c * w / tot for c in colw]
    hs = head_size or size + 0.5
    # alturas por contenido
    hts = []
    for r, row in enumerate(rows):
        hh = rowh
        for c, cell in enumerate(row):
            sz = hs if r == 0 else size
            th = text_h(cell, colw[c] - 0.16, sz, F_BODY if r else F_HEAD, r == 0 or (c == 0 and bold_first)) + 0.1
            hh = max(hh, th)
        hts.append(hh)
    if heights:
        hts = heights
    sid = s._nid(); nm = s._name(name or 'Tabla')
    grid = ''.join('<a:gridCol w="%d"/>' % E(c) for c in colw)
    trs = ''
    def bd(tag, col='1E1E1E', wd=1.0):
        return '<a:%s w="%d"><a:solidFill><a:srgbClr val="%s"/></a:solidFill></a:%s>' % (tag, int(wd * 12700), col, tag)
    for r, row in enumerate(rows):
        tcs = ''
        for c, cell in enumerate(row):
            if r == 0:
                f = head_fill; font = F_HEAD; sz = hs; col = head_color; b = True
            else:
                f = (first_col_fill if (c == 0 and first_col_fill) else fills[(r - 1) % len(fills)])
                font = F_BODY; sz = size; col = INK; b = (c == 0 and bold_first)
            paras = ''.join(para_xml(l, sz, font, col, 'l' if (c or r) else 'l', b, 1) for l in cell.split('\n'))
            tcs += ('<a:tc><a:txBody><a:bodyPr/><a:lstStyle/>%s</a:txBody><a:tcPr marL="%d" marR="%d" marT="%d" marB="%d" anchor="ctr">'
                    '%s%s%s%s<a:solidFill><a:srgbClr val="%s"/></a:solidFill></a:tcPr></a:tc>'
                    % (paras, E(0.08), E(0.08), E(0.04), E(0.04), bd('lnL'), bd('lnR'), bd('lnT'), bd('lnB'), f))
        trs += '<a:tr h="%d">%s</a:tr>' % (E(hts[r]), tcs)
    H = sum(hts)
    xml = ('<p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="%d" name="%s"/><p:cNvGraphicFramePr><a:graphicFrameLocks noGrp="1"/></p:cNvGraphicFramePr><p:nvPr/></p:nvGraphicFramePr>'
           '<p:xfrm><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></p:xfrm><a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/table">'
           '<a:tbl><a:tblPr firstRow="1" bandRow="1"/><a:tblGrid>%s</a:tblGrid>%s</a:tbl></a:graphicData></a:graphic></p:graphicFrame>'
           % (sid, esc(nm), E(x), E(y), E(w), E(H), grid, trs))
    s._append(xml)
    s._track(sid, x, y, w, H)
    if y + H > SH - 0.15:
        WARN.append('TABLA desborda: %s alto %.2f desde %.2f' % (nm, H, y))
    return sid, H


# ---------------------------------------------------------------- gráficos nativos
def chart_line(s, x, y, w, h, cats, series, colors, title_txt=None, ytitle=None, smooth=True, name='Gráfico',
               legend=True, size=12, stacked_area=False, ymax=None, xtitle=None, labels=False):
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.dml.color import RGBColor
    cd = CategoryChartData()
    cd.categories = cats
    for nm, vals in series:
        cd.add_series(nm, vals)
    ct = XL_CHART_TYPE.AREA_STACKED if stacked_area else XL_CHART_TYPE.LINE
    gf = s.ps.shapes.add_chart(ct, Emu(E(x)), Emu(E(y)), Emu(E(w)), Emu(E(h)), cd)
    gf.name = name
    ch = gf.chart
    ch.font.size = Pt(size); ch.font.name = F_BODY
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(size); ch.legend.font.name = F_BODY
    for i, ser in enumerate(ch.plots[0].series):
        if stacked_area:
            ser.format.fill.solid(); ser.format.fill.fore_color.rgb = RGBColor.from_string(colors[i])
            ser.format.line.color.rgb = RGBColor.from_string(INK); ser.format.line.width = Pt(0.75)
        else:
            ser.format.line.color.rgb = RGBColor.from_string(colors[i]); ser.format.line.width = Pt(3.5)
            ser.smooth = smooth
            from pptx.enum.chart import XL_MARKER_STYLE
            ser.marker.style = XL_MARKER_STYLE.NONE
    if labels:
        from pptx.enum.chart import XL_LABEL_POSITION
        pl = ch.plots[0]
        pl.has_data_labels = True
        pl.data_labels.font.size = Pt(size); pl.data_labels.font.bold = True; pl.data_labels.font.name = F_BODY
        pl.data_labels.position = XL_LABEL_POSITION.ABOVE
        pl.data_labels.show_value = True
    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RGBColor.from_string('D9B8A5')
    va.tick_labels.font.size = Pt(size - 1); va.tick_labels.font.name = F_BODY
    from pptx.enum.chart import XL_TICK_LABEL_POSITION
    if ymax:
        va.maximum_scale = ymax
    va.minimum_scale = 0
    va.tick_label_position = XL_TICK_LABEL_POSITION.NONE
    va.format.line.color.rgb = RGBColor.from_string(INK)
    if ytitle:
        va.has_title = True; va.axis_title.text_frame.text = ytitle
        r = va.axis_title.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(size); r.font.name = F_BODY; r.font.bold = True
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(size); ca.tick_labels.font.name = F_BODY
    ca.format.line.color.rgb = RGBColor.from_string(INK)
    if xtitle:
        ca.has_title = True; ca.axis_title.text_frame.text = xtitle
        r = ca.axis_title.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(size); r.font.name = F_BODY; r.font.bold = True
    if title_txt:
        ch.has_title = True; ch.chart_title.text_frame.text = title_txt
        r = ch.chart_title.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(size + 2); r.font.name = F_HEAD; r.font.bold = True
    else:
        ch.has_title = False
    # fondo del área del gráfico translúcido
    cs = ch._chartSpace
    from lxml import etree as ET
    sppr = ET.SubElement(cs, '{http://schemas.openxmlformats.org/drawingml/2006/chart}spPr')
    sppr.append(ET.fromstring('<a:solidFill xmlns:a="%s"><a:srgbClr val="FFFFFF"><a:alpha val="55000"/></a:srgbClr></a:solidFill>' % NS['a']))
    sppr.append(ET.fromstring('<a:ln xmlns:a="%s" w="15875"><a:solidFill><a:srgbClr val="1E1E1E"/></a:solidFill></a:ln>' % NS['a']))
    # mover spPr a la posición correcta (antes de txPr / externalData)
    order_ok = ['date1904', 'lang', 'roundedCorners', 'AlternateContent', 'clrMapOvr', 'pivotSource', 'protection', 'chart', 'spPr', 'txPr', 'externalData', 'printSettings', 'userShapes', 'extLst']
    kids = list(cs)
    kids.sort(key=lambda e: order_ok.index(ET.QName(e).localname) if ET.QName(e).localname in order_ok else 99)
    for k in kids:
        cs.remove(k)
    for k in kids:
        cs.append(k)
    s._track(gf.shape_id, x, y, w, h)
    return gf.shape_id


# ---------------------------------------------------------------- portada de unidad
def bullets(lines):
    return '\n'.join(lines)


def chain(s, x, y, items, dy=0.62, w=1.9, enzymes=None, rev=None, fill=YEL, size=13, ecol=None, name='Ruta', ew=None,
          horizontal=False, dx=None, h=0.4, side='r'):
    """ruta metabólica: fichas de metabolitos + flechas (doble punta = reversible; gruesa = irreversible) + enzimas.
    Devuelve [(id_metabolito, id_paso)] donde id_paso agrupa flecha y enzima."""
    res = []
    n = len(items)
    enzymes = enzymes or [None] * (n - 1)
    rev = rev or [True] * (n - 1)
    for i, it in enumerate(items):
        f = fill[i] if isinstance(fill, list) else fill
        if horizontal:
            cx = x + i * (dx or (w + 0.7)); cy = y
        else:
            cx = x; cy = y + i * dy
        stp = None
        if i:
            g = s.group_start('%s paso %d' % (name, i))
            if horizontal:
                p1 = (cx - (dx or (w + 0.7)) + w + 0.05, cy + h / 2); p2 = (cx - 0.05, cy + h / 2)
            else:
                p1 = (cx + w / 2, cy - dy + h + 0.03); p2 = (cx + w / 2, cy - 0.03)
            r = rev[i - 1]
            s.arrow([p1, p2], INK if r else (ecol or INK), 2.0 if r else 3.5, head='triangle' if r else None)
            e = enzymes[i - 1]
            if e:
                if horizontal:
                    ewid = ew or (text_w(plain(e), F_HEAD, True, 11.5) + 0.24)
                    enz(s, (p1[0] + p2[0]) / 2 - ewid / 2, cy - 0.42, e, w=ewid, size=11.5)
                else:
                    ewid = ew or (text_w(plain(e), F_HEAD, True, 11.5) + 0.24)
                    ex_ = cx + w / 2 + 0.12 if side == 'r' else cx + w / 2 - 0.12 - ewid
                    enz(s, ex_, (p1[1] + p2[1]) / 2 - 0.16, e, w=ewid, size=11.5, h=0.3)
            stp = s.group_end()
        m = s.shape(cx, cy, w, h, 'roundRect', f, (INK, 1.0), False, it, size, F_HEAD, INK, True, 'c', 'm', (0.03, 0, 0.03, 0),
                    name='%s %s' % (name, plain(it)[:18]), adj={'adj': 30000}, minsize=9)
        res.append((m, stp))
    return res


def chain_anim(s, res, per_click=1, start_new=True):
    first = True
    for k, (m, stp) in enumerate(res):
        effs = ([wipe(stp, 't')] if stp else []) + [fade(m, after=bool(stp))]
        if first or k % per_click == 0:
            s.step(*effs); first = False
        else:
            s.anim[-1] += [dict(e, after=True) if i == 0 else e for i, e in enumerate(effs)]
