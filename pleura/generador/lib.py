# -*- coding: utf-8 -*-
"""Motor de construcción: formas OOXML nativas, texto enriquecido,
animaciones (p:timing), transiciones Morph (p159:morph), fondos
anatómicos e incrustación de la fuente manuscrita."""
import re, os, copy
from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Pt
from PIL import ImageFont

EMU = 914400
SW, SH = 13.333, 7.5

NS = {
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
}
NSDECL = ' '.join('xmlns:%s="%s"' % kv for kv in NS.items())

# ---------------------------------------------------------------- tipografía
F_TITLE = 'Mali'
F_HEAD = 'Mali'
F_BODY = 'Mali'
F_COND = 'Mali'
F_MATH = 'Cambria Math'
F_SYM = 'Segoe UI Symbol'
HERE = os.path.dirname(os.path.abspath(__file__))
import fuentes as _fuentes
_fuentes.prepare()

# ---------------------------------------------------------------- paleta
INK = '1E1E1E'
PEACH = 'F7C6AA'      # fondo principal (durazno)
CORALBG = 'F3A9A2'    # fondo rosa coral
YELBG = 'FFF3A6'      # fondo amarillo claro
BLUEBG = 'D3E5F5'     # fondo celeste pastel
CREAMBG = 'FBE4D2'
TAG = 'FFC000'        # amarillo intenso (etiquetas)
SALMON = 'F4B183'
LSALMON = 'FAD3BC'
YEL = 'FFE07A'
CREAM = 'FFF4D6'
RED = 'D7261E'
PINK = 'F9C6D0'
CORAL = 'F28B82'
BLUE = '9CC3E8'
LBLUE = 'DCEBF8'
MINT = 'CDE8C4'
LILAC = 'DCCDF0'
ORANGE = 'F28C28'
WHITE = 'FFFFFF'
# código funcional
ACT = '1E8C3A'   # activación
INH = 'C00000'   # inhibición
TRN = '2563B0'   # transporte
ENE = 'FFC000'   # intermediarios energéticos
NIT = '0F7C7C'   # nitrógeno

# colores moleculares coherentes en toda la presentación
MOL = {
    'glc': ('3B82C4', WHITE), 'lac': ('8E44AD', WHITE), 'ala': ('2E9E5B', WHITE),
    'ag': ('F2B90F', INK), 'gly': ('F28C28', INK), 'cc': ('C2185B', WHITE),
    'nh4': ('0F7C7C', WHITE), 'urea': ('7F8C8D', WHITE), 'gln': ('16A085', WHITE),
    'aa': ('27AE60', WHITE), 'ins': ('1F4E9A', WHITE), 'gcg': ('D35400', WHITE),
    'atp': ('FFC000', INK), 'ca': ('E74C3C', WHITE), 'k': ('6C3483', WHITE),
    'tag': ('F7DC6F', INK), 'chol': ('F39C12', INK), 'o2': ('5DADE2', INK),
    'co2': ('95A5A6', INK), 'urate': ('A04000', WHITE), 'fru': ('E67E22', WHITE),
    'acoa': ('B7950B', WHITE), 'mal': ('7D3C98', WHITE),
}

# ---------------------------------------------------------------- medición de texto
_FONTCACHE = {}
FONTFILES = {('Mali', False): 'Mali-Regular.ttf', ('Mali', True): 'Mali-Bold.ttf'}


def _font(face, bold, size):
    key = (face, bold, size)
    if key not in _FONTCACHE:
        fn = FONTFILES.get((face, bold), FONTFILES.get((face, False), 'Mali-Regular.ttf'))
        _FONTCACHE[key] = ImageFont.truetype(os.path.join(HERE, 'fuentes', fn), int(size * 10))
    return _FONTCACHE[key]


def text_w(s, face, bold, size):
    """ancho en pulgadas medido con la fuente real (4 % de holgura)"""
    f = _font(face, bold, size)
    return f.getlength(s) / 10 / 72 * 1.04


# Interlineado EXACTO (en puntos) = LINE × tamaño × factor. Se escribe como <a:spcPts>, que
# PowerPoint (Windows/Mac), Google Slides y LibreOffice respetan igual, sin depender de las
# métricas verticales de la fuente. La estimación de alto de texto usa exactamente el mismo valor.
LINE = 1.22


WARN = []

# ---------------------------------------------------------------- texto enriquecido
COLTAG = {'g': ACT, 'r': INH, 'b': TRN, 'o': 'C55A11', 'p': '7030A0', 'w': WHITE, 'k': INK,
          'n': NIT, 'y': 'B8860B'}
TOK = re.compile(r'(\*\*|\^\{[^}]*\}|_\{[^}]*\}|\[[a-z]\]|\[/\]|~~)')


def esc(t):
    return t.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')


def parse_runs(text):
    """devuelve [(texto, bold, italic, color, baseline)]"""
    out = []
    bold = False; ital = False; col = None
    for tok in TOK.split(text):
        if not tok:
            continue
        if tok == '**':
            bold = not bold
        elif tok == '~~':
            ital = not ital
        elif tok.startswith('^{'):
            out.append((tok[2:-1], bold, ital, col, 30000))
        elif tok.startswith('_{'):
            out.append((tok[2:-1], bold, ital, col, -25000))
        elif re.fullmatch(r'\[[a-z]\]', tok):
            col = COLTAG[tok[1]]; bold = True
        elif tok == '[/]':
            col = None; bold = False
        else:
            out.append((tok, bold, ital, col, 0))
    return out


def plain(text):
    return ''.join(r[0] for r in parse_runs(text))


def _cmap():
    global _CMAP
    if _CMAP is None:
        from fontTools.ttLib import TTFont
        _CMAP = set(TTFont(os.path.join(HERE, 'fuentes', 'Mali-Regular.ttf')).getBestCmap())
    return _CMAP


_CMAP = None


def _split_glyphs(t):
    """separa caracteres ausentes en Mali (σ, µ, ✓…) para asignarles una fuente de respaldo"""
    out = []
    cm = _cmap()
    for ch in t:
        miss = (ord(ch) not in cm and ch not in '\n') or ch in '•'
        fb = ('Arial' if ch == '•' else F_MATH if ch in 'αβγδεζηθικλμνξοπρστυφχψωµΔΣΩ' else F_SYM) if miss else None
        if out and out[-1][1] == fb:
            out[-1][0] += ch
        else:
            out.append([ch, fb])
    return out


def runs_xml(text, size, font, color, bold=False, ital=False, reff=''):
    xs = []
    for t0, b, i, c, base in parse_runs(text):
      for t, fb in _split_glyphs(t0):
        f = fb or font
        attrs = 'lang="es-AR" sz="%d"' % int(size * 100)
        if b or bold:
            attrs += ' b="1"'
        if i or ital:
            attrs += ' i="1"'
        if base:
            attrs += ' baseline="%d"' % base
        attrs += ' dirty="0"'
        xs.append('<a:r><a:rPr %s><a:solidFill><a:srgbClr val="%s"/></a:solidFill>'
                  '%s<a:latin typeface="%s"/><a:ea typeface="%s"/><a:cs typeface="%s"/><a:sym typeface="%s"/></a:rPr><a:t>%s</a:t></a:r>'
                  % (attrs, c or color, reff, f, f, f, f, esc(t)))
    return ''.join(xs)


def para_xml(line, size, font, color, align='l', bold=False, spc_after=0, ital=False, lnspc=None, reff=''):
    bullet = ''
    indent = ''
    if line.startswith('- '):
        line = line[2:]
        bullet = '<a:buClr><a:srgbClr val="%s"/></a:buClr><a:buFont typeface="Arial"/><a:buChar char="&#8226;"/>' % color
        indent = ' marL="%d" indent="-%d"' % (int(size * 0.75 / 72 * EMU), int(size * 0.75 / 72 * EMU))
    elif line.startswith('+ '):
        line = line[2:]
        bullet = '<a:buClr><a:srgbClr val="%s"/></a:buClr><a:buFont typeface="Segoe UI Symbol"/><a:buChar char="&#10003;"/>' % ACT
        indent = ' marL="%d" indent="-%d"' % (int(size * 0.95 / 72 * EMU), int(size * 0.95 / 72 * EMU))
    elif line.startswith('>> '):
        line = line[3:]
        bullet = '<a:buClr><a:srgbClr val="%s"/></a:buClr><a:buFont typeface="Segoe UI Symbol"/><a:buChar char="&#8594;"/>' % 'C55A11'
        indent = ' marL="%d" indent="-%d"' % (int(size * 1.0 / 72 * EMU), int(size * 1.0 / 72 * EMU))
    elif line.startswith('xx '):
        line = line[3:]
        bullet = '<a:buClr><a:srgbClr val="%s"/></a:buClr><a:buFont typeface="Segoe UI Symbol"/><a:buChar char="&#10007;"/>' % INH
        indent = ' marL="%d" indent="-%d"' % (int(size * 0.95 / 72 * EMU), int(size * 0.95 / 72 * EMU))
    else:
        bullet = '<a:buNone/>'
    sp = '<a:spcAft><a:spcPts val="%d"/></a:spcAft>' % int(spc_after * 100) if spc_after else ''
    ls = '<a:lnSpc><a:spcPts val="%d"/></a:lnSpc>' % int(round(size * LINE * (lnspc or 1.0) * 100))
    algn = {'l': 'l', 'c': 'ctr', 'r': 'r', 'j': 'just'}[align]
    end = '<a:endParaRPr lang="es-AR" sz="%d" dirty="0"><a:latin typeface="%s"/></a:endParaRPr>' % (int(size * 100), font)
    return '<a:p><a:pPr algn="%s"%s>%s%s%s</a:pPr>%s%s</a:p>' % (
        algn, indent, ls, sp, bullet, runs_xml(line, size, font, color, bold, ital, reff) if line else '', end)


def _words(line, bold):
    """palabras de una línea como listas de (carácter, negrita)"""
    words = [[]]
    for t, b, i, c, base in parse_runs(line):
        for ch in t:
            if ch == ' ':
                words.append([])
            else:
                words[-1].append((ch, b or bold))
    return words


def _word_w(word, font, size):
    """ancho de una palabra midiendo cada tramo con su peso (regular / negrita)"""
    w = 0.0; cur = ''; cb = None
    for ch, b in word:
        if cb is None or b == cb:
            cur += ch; cb = b
        else:
            w += text_w(cur, font, cb, size); cur = ch; cb = b
    if cur:
        w += text_w(cur, font, cb, size)
    return w


def layout(text, w_in, size, font, bold):
    """(líneas totales, hay_palabra_más_ancha_que_el_recuadro)"""
    n = 0; over = False
    sp = text_w(' ', font, False, size)
    for line in text.split('\n'):
        ind = 0
        if line[:2] == '- ':
            line = line[2:]; ind = size * 0.75 / 72
        elif line[:2] == '+ ':
            line = line[2:]; ind = size * 0.95 / 72
        avail = w_in - ind
        lines = 1; x = 0.0
        for wd in _words(line, bold):
            ww = _word_w(wd, font, size)
            if ww > avail:
                over = True
            if x == 0:
                x = ww
            elif x + sp + ww <= avail:
                x += sp + ww
            else:
                lines += 1; x = ww
        n += lines
    return n, over


def wrap_lines(text, w_in, size, font, bold):
    return layout(text, w_in, size, font, bold)[0]


def text_h(text, w_in, size, font, bold, spc_after=0, lnspc=1.0):
    nl = wrap_lines(text, w_in, size, font, bold)
    npar = text.count('\n') + 1
    return nl * size * LINE * (lnspc or 1.0) / 72 + (npar - 1) * spc_after / 72


def fits(text, w_in, h_in, size, font, bold=False, spc_after=0, lnspc=1.0):
    nl, over = layout(text, w_in, size, font, bold)
    npar = text.count('\n') + 1
    h = nl * size * LINE * (lnspc or 1.0) / 72 + (npar - 1) * spc_after / 72
    return (h <= h_in) and not over


def fit_size(text, w_in, h_in, size, font, bold=False, minsize=11, spc_after=0, lnspc=1.0, label=''):
    s = size
    while s > minsize and not fits(text, w_in, h_in, s, font, bold, spc_after * s / size, lnspc):
        s -= 0.5
    nl, over = layout(text, w_in, s, font, bold)
    if text_h(text, w_in, s, font, bold, spc_after * s / size, lnspc) > h_in * 1.02:
        WARN.append('DESBORDE (%s): %.1fpt en %.2fx%.2f -> %r' % (label, s, w_in, h_in, plain(text)[:60]))
    if over:
        WARN.append('PALABRA ANCHA (%s): %.1fpt en %.2f -> %r' % (label, s, w_in, plain(text)[:60]))
    return s


# ---------------------------------------------------------------- geometría y estilo
def E(v):
    return int(round(v * EMU))


def fill_xml(fill):
    if fill is None:
        return '<a:noFill/>'
    if isinstance(fill, str):
        return '<a:solidFill><a:srgbClr val="%s"/></a:solidFill>' % fill
    kind = fill[0]
    if kind == 'a':  # ('a', color, alpha%)
        return '<a:solidFill><a:srgbClr val="%s"><a:alpha val="%d"/></a:srgbClr></a:solidFill>' % (fill[1], fill[2] * 1000)
    stops = ''.join('<a:gs pos="%d"><a:srgbClr val="%s"%s</a:gs>' % (
        p * 1000, c, ('><a:alpha val="%d"/></a:srgbClr>' % (al * 1000)) if al is not None else '/>')
        for (p, c, *rest) in fill[1] for al in [rest[0] if rest else None])
    if kind == 'lin':
        ang = fill[2] if len(fill) > 2 else 90
        return '<a:gradFill rotWithShape="1"><a:gsLst>%s</a:gsLst><a:lin ang="%d" scaled="0"/></a:gradFill>' % (stops, ang * 60000)
    if kind == 'rad':
        l, t = (fill[2] if len(fill) > 2 else (35, 30))
        return ('<a:gradFill rotWithShape="1"><a:gsLst>%s</a:gsLst><a:path path="circle">'
                '<a:fillToRect l="%d" t="%d" r="%d" b="%d"/></a:path></a:gradFill>' % (stops, l * 1000, t * 1000, (100 - l) * 1000, (100 - t) * 1000))
    raise ValueError(fill)


def ln_xml(line):
    """line: None | color | (color, pt) | (color, pt, dash, head, tail)"""
    if line is None:
        return '<a:ln><a:noFill/></a:ln>'
    if isinstance(line, str):
        line = (line, 1.25)
    col, w = line[0], line[1]
    dash = line[2] if len(line) > 2 else None
    head = line[3] if len(line) > 3 else None
    tail = line[4] if len(line) > 4 else None
    x = '<a:ln w="%d" cap="rnd">' % int(w * 12700)
    x += fill_xml(col) if col else '<a:noFill/>'
    if dash:
        x += '<a:prstDash val="%s"/>' % dash
    x += '<a:round/>'
    if head:
        x += '<a:headEnd type="%s" w="med" len="med"/>' % head
    if tail:
        x += '<a:tailEnd type="%s" w="med" len="med"/>' % tail
    return x + '</a:ln>'


def shadow_xml(sh):
    if not sh:
        return ''
    if isinstance(sh, str):
        return sh
    if sh is True:
        sh = (4, 3, 35)
    blur, dist, alpha = sh
    return ('<a:effectLst><a:outerShdw blurRad="%d" dist="%d" dir="2700000" algn="tl" rotWithShape="0">'
            '<a:srgbClr val="000000"><a:alpha val="%d"/></a:srgbClr></a:outerShdw></a:effectLst>' % (blur * 12700, dist * 12700, alpha * 1000))


PATHTOK = re.compile(r'[MLCQZAmlcqz]|-?\d+(?:\.\d+)?')


def path_xml(d, W=1000, H=1000, fill=True, stroke=True):
    toks = PATHTOK.findall(d)
    i = 0; out = []
    cmd = None
    def num():
        nonlocal i
        v = float(toks[i]); i += 1
        return int(round(v))
    while i < len(toks):
        t = toks[i]
        if t in 'MLCQZ':
            cmd = t; i += 1
            if t == 'Z':
                out.append('<a:close/>'); continue
        if cmd == 'M':
            out.append('<a:moveTo><a:pt x="%d" y="%d"/></a:moveTo>' % (num(), num())); cmd = 'L'
        elif cmd == 'L':
            out.append('<a:lnTo><a:pt x="%d" y="%d"/></a:lnTo>' % (num(), num()))
        elif cmd == 'C':
            pts = [(num(), num()) for _ in range(3)]
            out.append('<a:cubicBezTo>%s</a:cubicBezTo>' % ''.join('<a:pt x="%d" y="%d"/>' % p for p in pts))
        elif cmd == 'Q':
            pts = [(num(), num()) for _ in range(2)]
            out.append('<a:quadBezTo>%s</a:quadBezTo>' % ''.join('<a:pt x="%d" y="%d"/>' % p for p in pts))
        else:
            raise ValueError(d)
    attrs = 'w="%d" h="%d"' % (W, H)
    if not fill:
        attrs += ' fill="none"'
    if not stroke:
        attrs += ' stroke="0"'
    return '<a:path %s>%s</a:path>' % (attrs, ''.join(out))


def custgeom(paths):
    return ('<a:custGeom><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/><a:rect l="0" t="0" r="r" b="b"/>'
            '<a:pathLst>%s</a:pathLst></a:custGeom>' % ''.join(paths))


def prst(name, adj=None):
    av = ''
    if adj:
        av = ''.join('<a:gd name="%s" fmla="val %d"/>' % (k, v) for k, v in adj.items())
    return '<a:prstGeom prst="%s"><a:avLst>%s</a:avLst></a:prstGeom>' % (name, av)


# ---------------------------------------------------------------- slide
class Slide:
    def __init__(self, deck, pslide, bg):
        self.deck = deck
        self.ps = pslide
        self.tree = pslide.shapes._spTree
        self.names = set()
        self.groups = []      # pila de grupos abiertos
        self.anim = []        # lista de pasos (clic) -> lista de efectos
        self.morph = False
        self.notes = ''
        self.bbox = {}
        self.set_bg(bg)

    # -- ids y nombres
    def _nid(self):
        ids = [int(x) for x in self.tree.xpath('.//p:cNvPr/@id')]
        if not hasattr(self, '_cid'):
            self._cid = max(ids + [1])
        self._cid = max(self._cid, max(ids + [1])) + 1
        return self._cid

    def _name(self, name, kind='Forma'):
        if not name:
            name = '%s %d' % (kind, len(self.names) + 1)
        base = name; k = 2
        while name in self.names:
            name = '%s_%d' % (base, k); k += 1
        self.names.add(name)
        return name

    def set_bg(self, bg):
        """bg: color, relleno o ruta de imagen PNG (fondo anatómico desenfocado)"""
        cSld = self.ps._element.find('{%s}cSld' % NS['p'])
        old = cSld.find('{%s}bg' % NS['p'])
        if old is not None:
            cSld.remove(old)
        if isinstance(bg, str) and bg.endswith('.png'):
            _, rid = self.ps.part.get_or_add_image_part(bg)
            f = ('<a:blipFill dpi="0" rotWithShape="1"><a:blip r:embed="%s"/><a:srcRect/><a:stretch><a:fillRect/></a:stretch></a:blipFill>' % rid)
        else:
            f = fill_xml(bg)
        x = '<p:bg %s><p:bgPr>%s<a:effectLst/></p:bgPr></p:bg>' % (NSDECL, f)
        cSld.insert(0, etree.fromstring(x))

    def _append(self, xml):
        el = etree.fromstring('<p:wrap %s>%s</p:wrap>' % (NSDECL, xml))[0]
        if self.groups:
            self.groups[-1]['el'].append(el)
        else:
            self.tree.append(el)
        return el

    def _track(self, sid, x, y, w, h):
        self.bbox[sid] = (x, y, w, h)
        for g in self.groups:
            bx = g['bb']
            if bx is None:
                g['bb'] = [x, y, x + w, y + h]
            else:
                bx[0] = min(bx[0], x); bx[1] = min(bx[1], y); bx[2] = max(bx[2], x + w); bx[3] = max(bx[3], y + h)

    # -- forma genérica
    def shape(self, x, y, w, h, geom='rect', fill=None, line=INK, shadow=False, text=None, size=14,
              font=F_BODY, color=INK, bold=False, align='l', anchor='t', ins=(0.08, 0.05, 0.08, 0.05),
              rot=0, name=None, adj=None, fit=True, minsize=11, spc=3, wrap=True, flipH=False, flipV=False,
              ital=False, lnspc=None, txt_rot=None, alt=None, reff=''):
        sid = self._nid(); nm = self._name(name)
        if isinstance(geom, str) and not geom.startswith('<'):
            g = prst(geom, adj)
        else:
            g = geom
        tb = ''
        if text is not None:
            iw = w - ins[0] - ins[2]; ih = h - ins[1] - ins[3]
            if fit and wrap:
                size = fit_size(text, iw, ih, size, font, bold, minsize, spc, lnspc or 1.0, label=nm)
            paras = ''.join(para_xml(l, size, font, color, align, bold, spc, ital, lnspc, reff) for l in text.split('\n'))
            vert = ' vert="vert270"' if txt_rot == 270 else ''
            tb = ('<p:txBody><a:bodyPr wrap="%s" lIns="%d" tIns="%d" rIns="%d" bIns="%d" anchor="%s" rtlCol="0"%s>'
                  '<a:noAutofit/></a:bodyPr><a:lstStyle/>%s</p:txBody>' % (
                      'square' if wrap else 'none', E(ins[0]), E(ins[1]), E(ins[2]), E(ins[3]),
                      {'t': 't', 'm': 'ctr', 'b': 'b'}[anchor], vert, paras))
        rotx = (' rot="%d"' % int(rot * 60000)) if rot else ''
        fl = (' flipH="1"' if flipH else '') + (' flipV="1"' if flipV else '')
        txbox = ' txBox="1"' if (fill is None and line is None and text is not None) else ''
        descr = ' descr="%s"' % esc(alt).replace('"', "'") if alt else ''
        xml = ('<p:sp><p:nvSpPr><p:cNvPr id="%d" name="%s"%s/><p:cNvSpPr%s/><p:nvPr/></p:nvSpPr>'
               '<p:spPr><a:xfrm%s%s><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm>%s%s%s%s</p:spPr>%s</p:sp>' % (
                   sid, esc(nm), descr, txbox, rotx, fl, E(x), E(y), max(E(w), 1), max(E(h), 1), g, fill_xml(fill), ln_xml(line),
                   shadow_xml(shadow), tb))
        self._append(xml)
        self._track(sid, x, y, w, h)
        return sid

    def free(self, x, y, w, h, d, fill=None, line=INK, shadow=False, name=None, W=1000, H=1000, rot=0, **kw):
        """forma libre: d en coordenadas 0..W x 0..H"""
        paths = d if isinstance(d, list) else [d]
        px = []
        for p in paths:
            if isinstance(p, tuple):  # (d, fillflag, strokeflag)
                px.append(path_xml(p[0], W, H, p[1], p[2]))
            else:
                px.append(path_xml(p, W, H, fill is not None, True))
        return self.shape(x, y, w, h, custgeom(px), fill, line, shadow, name=name, rot=rot, **kw)

    # -- grupos
    def group_start(self, name=None):
        sid = self._nid(); nm = self._name(name, 'Grupo')
        xml = ('<p:grpSp><p:nvGrpSpPr><p:cNvPr id="%d" name="%s"/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
               '<p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/><a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr></p:grpSp>' % (sid, esc(nm)))
        el = self._append(xml)
        self.groups.append({'el': el, 'bb': None, 'id': sid})
        return sid

    def group_end(self):
        g = self.groups.pop()
        bb = g['bb'] or [0, 0, 0.01, 0.01]
        xf = g['el'].find('.//{%s}xfrm' % NS['a'])
        for tag, vals in (('off', (bb[0], bb[1])), ('chOff', (bb[0], bb[1]))):
            e = xf.find('{%s}%s' % (NS['a'], tag)); e.set('x', str(E(vals[0]))); e.set('y', str(E(vals[1])))
        for tag in ('ext', 'chExt'):
            e = xf.find('{%s}%s' % (NS['a'], tag)); e.set('cx', str(max(E(bb[2] - bb[0]), 1))); e.set('cy', str(max(E(bb[3] - bb[1]), 1)))
        self.bbox[g['id']] = (bb[0], bb[1], bb[2] - bb[0], bb[3] - bb[1])
        if self.groups:  # propaga
            pb = self.groups[-1]['bb']
            if pb is None:
                self.groups[-1]['bb'] = list(bb)
            else:
                pb[0] = min(pb[0], bb[0]); pb[1] = min(pb[1], bb[1]); pb[2] = max(pb[2], bb[2]); pb[3] = max(pb[3], bb[3])
        return g['id']

    # -- texto suelto
    def text(self, x, y, w, h, text, size=14, font=F_BODY, color=INK, bold=False, align='l', anchor='t', name=None, **kw):
        kw.setdefault('ins', (0, 0, 0, 0))
        return self.shape(x, y, w, h, 'rect', None, None, False, text, size, font, color, bold, align, anchor, name=name, **kw)

    # -- flechas (trazos abiertos con punta)
    def arrow(self, pts, color=INK, w=2.25, head=None, tail='triangle', dash=None, curve=False, name=None, shadow=False):
        """pts: lista de (x,y) en pulgadas; curve=True usa Bézier cuadrática por punto medio"""
        xs = [p[0] for p in pts]; ys = [p[1] for p in pts]
        x0, y0 = min(xs), min(ys)
        bw = max(max(xs) - x0, 0.01); bh = max(max(ys) - y0, 0.01)
        W = max(int(bw * 1000), 1); H = max(int(bh * 1000), 1)
        def c(p):
            return '%d %d' % (int((p[0] - x0) * 1000), int((p[1] - y0) * 1000))
        if curve and len(pts) == 3:
            d = 'M %s Q %s %s' % (c(pts[0]), c(pts[1]), c(pts[2]))
        elif curve and len(pts) == 4:
            d = 'M %s C %s %s %s' % (c(pts[0]), c(pts[1]), c(pts[2]), c(pts[3]))
        else:
            d = 'M ' + ' L '.join(c(p) for p in pts)
        return self.shape(x0, y0, bw, bh, custgeom([path_xml(d, W, H, False, True)]), None,
                          (color, w, dash, head, tail), shadow, name=name or self._name(None, 'Flecha'))

    def inhib(self, pts, color=INH, w=2.5, name=None, blen=0.22):
        """flecha de inhibición: trazo + barra perpendicular (⊣)"""
        import math
        g = self.group_start(name or 'Inhibición')
        self.arrow(pts, color, w, tail=None)
        (x1, y1), (x2, y2) = pts[-2], pts[-1]
        L = math.hypot(x2 - x1, y2 - y1) or 1
        nx, ny = -(y2 - y1) / L * blen / 2, (x2 - x1) / L * blen / 2
        self.arrow([(x2 - nx, y2 - ny), (x2 + nx, y2 + ny)], color, w * 1.6, tail=None)
        self.group_end()
        return g

    # -- animación
    def step(self, *effects):
        self.anim.append(list(effects))

    # -- notas
    def note(self, t):
        self.notes += (('\n\n' if self.notes else '') + t)


# ---------------------------------------------------------------- efectos
def fx(kind, sid, after=False, delay=0, dur=None, **kw):
    d = {'kind': kind, 'sid': sid, 'after': after, 'delay': delay}
    d['dur'] = dur or {'appear': 1, 'fade': 500, 'wipe': 600, 'zoom': 500, 'path': 1500, 'out': 400,
                       'pulse': 300, 'spin': 1000, 'outnow': 1, 'grow': 600, 'float': 600}[kind]
    d.update(kw)
    return d


def appear(s, after=False, delay=0):
    return fx('appear', s, after, delay)


def fade(s, after=False, delay=0, dur=None):
    return fx('fade', s, after, delay, dur)


def wipe(s, dir='l', after=False, delay=0, dur=None):
    return fx('wipe', s, after, delay, dur, dir=dir)


def zoom(s, after=False, delay=0, dur=None):
    return fx('zoom', s, after, delay, dur)


def path(s, dx, dy, after=False, delay=0, dur=None, pts=None):
    """desplazamiento relativo en pulgadas (o lista de puntos intermedios)"""
    return fx('path', s, after, delay, dur, dx=dx, dy=dy, pts=pts)


def out(s, after=False, delay=0, dur=None):
    return fx('out', s, after, delay, dur)


def pulse(s, after=False, delay=0, dur=None):
    return fx('pulse', s, after, delay, dur)


def spin(s, after=False, delay=0, dur=None):
    return fx('spin', s, after, delay, dur)


class TimingBuilder:
    def __init__(self):
        self.n = 2

    def nid(self):
        self.n += 1
        return self.n

    def tgt(self, sid):
        return '<p:tgtEl><p:spTgt spid="%d"/></p:tgtEl>' % sid

    def set_vis(self, sid, val, delay=0):
        return ('<p:set><p:cBhvr><p:cTn id="%d" dur="1" fill="hold"><p:stCondLst><p:cond delay="%d"/></p:stCondLst></p:cTn>%s'
                '<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="%s"/></p:to></p:set>'
                % (self.nid(), delay, self.tgt(sid), val))

    def effect(self, e, node, grp):
        k = e['kind']; sid = e['sid']; dur = e['dur']
        cls, pid, sub = {'appear': ('entr', 1, 0), 'fade': ('entr', 10, 0), 'wipe': ('entr', 22, 0), 'zoom': ('entr', 53, 16),
                         'path': ('path', 0, 0), 'out': ('exit', 10, 0), 'outnow': ('exit', 1, 0), 'pulse': ('emph', 6, 0),
                         'spin': ('emph', 8, 0), 'grow': ('entr', 53, 16), 'float': ('entr', 42, 0)}[k]
        if k == 'wipe':
            sub = {'t': 1, 'r': 2, 'b': 4, 'l': 8}[e['dir']]
        cid = self.nid()
        body = ''
        if k == 'appear':
            body = self.set_vis(sid, 'visible')
        elif k == 'fade':
            body = self.set_vis(sid, 'visible') + (
                '<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="%d" dur="%d"/>%s</p:cBhvr></p:animEffect>'
                % (self.nid(), dur, self.tgt(sid)))
        elif k == 'wipe':
            dirn = {'t': 'up', 'b': 'down', 'l': 'left', 'r': 'right'}[e['dir']]
            # dir = lado desde el que comienza el barrido
            body = self.set_vis(sid, 'visible') + (
                '<p:animEffect transition="in" filter="wipe(%s)"><p:cBhvr><p:cTn id="%d" dur="%d"/>%s</p:cBhvr></p:animEffect>'
                % (dirn, self.nid(), dur, self.tgt(sid)))
        elif k in ('zoom', 'grow'):
            body = self.set_vis(sid, 'visible') + (
                '<p:anim calcmode="lin" valueType="num"><p:cBhvr><p:cTn id="%d" dur="%d" fill="hold"/>%s<p:attrNameLst><p:attrName>ppt_w</p:attrName></p:attrNameLst></p:cBhvr>'
                '<p:tavLst><p:tav tm="0"><p:val><p:fltVal val="0"/></p:val></p:tav><p:tav tm="100000"><p:val><p:strVal val="#ppt_w"/></p:val></p:tav></p:tavLst></p:anim>'
                '<p:anim calcmode="lin" valueType="num"><p:cBhvr><p:cTn id="%d" dur="%d" fill="hold"/>%s<p:attrNameLst><p:attrName>ppt_h</p:attrName></p:attrNameLst></p:cBhvr>'
                '<p:tavLst><p:tav tm="0"><p:val><p:fltVal val="0"/></p:val></p:tav><p:tav tm="100000"><p:val><p:strVal val="#ppt_h"/></p:val></p:tav></p:tavLst></p:anim>'
                '<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="%d" dur="%d"/>%s</p:cBhvr></p:animEffect>'
                % (self.nid(), dur, self.tgt(sid), self.nid(), dur, self.tgt(sid), self.nid(), dur, self.tgt(sid)))
        elif k == 'float':
            body = self.set_vis(sid, 'visible') + (
                '<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="%d" dur="%d"/>%s</p:cBhvr></p:animEffect>'
                '<p:anim calcmode="lin" valueType="num"><p:cBhvr><p:cTn id="%d" dur="%d" decel="100000"/>%s<p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr>'
                '<p:tavLst><p:tav tm="0"><p:val><p:strVal val="#ppt_y+.1"/></p:val></p:tav><p:tav tm="100000"><p:val><p:strVal val="#ppt_y"/></p:val></p:tav></p:tavLst></p:anim>'
                % (self.nid(), dur, self.tgt(sid), self.nid(), dur, self.tgt(sid)))
        elif k == 'out':
            body = ('<p:animEffect transition="out" filter="fade"><p:cBhvr><p:cTn id="%d" dur="%d"/>%s</p:cBhvr></p:animEffect>'
                    % (self.nid(), dur, self.tgt(sid))) + self.set_vis(sid, 'hidden', dur - 1)
        elif k == 'outnow':
            body = self.set_vis(sid, 'hidden')
        elif k == 'pulse':
            body = ('<p:animScale><p:cBhvr><p:cTn id="%d" dur="%d" autoRev="1" fill="hold"/>%s</p:cBhvr>'
                    '<p:by x="112000" y="112000"/></p:animScale>' % (self.nid(), dur, self.tgt(sid)))
        elif k == 'spin':
            body = ('<p:animRot by="21600000"><p:cBhvr><p:cTn id="%d" dur="%d" fill="hold"/>%s'
                    '<p:attrNameLst><p:attrName>r</p:attrName></p:attrNameLst></p:cBhvr></p:animRot>' % (self.nid(), dur, self.tgt(sid)))
        elif k == 'path':
            dx = e['dx'] / SW; dy = e['dy'] / SH
            if e.get('pts'):
                segs = ' '.join('L %.4f %.4f' % (px / SW, py / SH) for px, py in e['pts'])
                pth = 'M 0 0 %s L %.4f %.4f E' % (segs, dx, dy)
            else:
                pth = 'M 0 0 L %.4f %.4f E' % (dx, dy)
            body = ('<p:animMotion origin="layout" path="%s" pathEditMode="relative" ptsTypes="">'
                    '<p:cBhvr><p:cTn id="%d" dur="%d" fill="hold"/>%s<p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr>'
                    '<p:rCtr x="%d" y="%d"/></p:animMotion>' % (pth, self.nid(), dur, self.tgt(sid), int(dx * 50000), int(dy * 50000)))
        acc = ' accel="20000" decel="20000"' if k == 'path' else ''
        return ('<p:par><p:cTn id="%d" presetID="%d" presetClass="%s" presetSubtype="%d"%s fill="hold" grpId="%d" nodeType="%s">'
                '<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>%s</p:childTnLst></p:cTn></p:par>'
                % (cid, pid, cls, sub, acc, grp, node, body))

    def build(self, steps, spkinds):
        clicks = []
        grpids = {}
        used = {}
        for st in steps:
            if not st:
                continue
            # agrupar por tiempo de inicio
            prev_end = 0; slots = []  # (start, [effects])
            for i, e in enumerate(st):
                if i == 0:
                    start = e['delay']; newslot = True
                elif e['after']:
                    start = prev_end + e['delay']; newslot = True
                else:
                    start = slots[-1][0] + e['delay']; newslot = start != slots[-1][0]
                if newslot:
                    slots.append((start, [e]))
                else:
                    slots[-1][1].append(e)
                prev_end = max(prev_end, start + e['dur'])
            inner = []
            for j, (start, effs) in enumerate(slots):
                ex = []
                for k, e in enumerate(effs):
                    if j == 0 and k == 0:
                        node = 'clickEffect'
                    elif k == 0:
                        node = 'afterEffect'
                    else:
                        node = 'withEffect'
                    key = (e['sid'], e['kind'] in ('out', 'outnow'))
                    g = 0
                    used.setdefault(e['sid'], set()).add(g)
                    ex.append(self.effect(e, node, g))
                inner.append('<p:par><p:cTn id="%d" fill="hold"><p:stCondLst><p:cond delay="%d"/></p:stCondLst><p:childTnLst>%s</p:childTnLst></p:cTn></p:par>'
                             % (self.nid(), start, ''.join(ex)))
            clicks.append('<p:par><p:cTn id="%d" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>%s</p:childTnLst></p:cTn></p:par>'
                          % (self.nid(), ''.join(inner)))
        if not clicks:
            return ''
        bld = ''.join('<p:bldP spid="%d" grpId="0" animBg="1"/>' % sid for sid in used if spkinds.get(sid) == 'sp')
        return ('<p:timing %s><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
                '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>%s</p:childTnLst></p:cTn>'
                '<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
                '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
                '</p:childTnLst></p:cTn></p:par></p:tnLst>%s</p:timing>'
                % (NSDECL, ''.join(clicks), ('<p:bldLst>%s</p:bldLst>' % bld) if bld else ''))


MORPH = ('<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
         '<mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" '
         'xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" Requires="p159">'
         '<p:transition %s spd="slow" p14:dur="1600"><p159:morph option="byObject"/></p:transition></mc:Choice>'
         '<mc:Fallback><p:transition %s spd="slow"><p:fade/></p:transition></mc:Fallback></mc:AlternateContent>')
FADE = '<p:transition %s spd="med"><p:fade/></p:transition>'


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Emu(E(SW)); self.prs.slide_height = Emu(E(SH))
        self.layout = self.prs.slide_layouts[6]
        self.slides = []
        self.stats = {'anim_slides': 0, 'effects': 0, 'morph': 0}

    def new(self, bg=PEACH):
        ps = self.prs.slides.add_slide(self.layout)
        s = Slide(self, ps, bg)
        self.slides.append(s)
        return s

    def finalize(self):
        for s in self.slides:
            sld = s.ps._element
            # títulos y numeración siempre al frente
            for el in list(s.tree):
                c = el.find('.//{%s}cNvPr' % NS['p'])
                if c is not None and c.get('name') in ('Título', 'Etiqueta sección', 'Número', '!!Paso'):
                    s.tree.remove(el); s.tree.append(el)
            clr = sld.find('{%s}clrMapOvr' % NS['p'])
            idx = list(sld).index(clr) + 1 if clr is not None else len(sld)
            if s.morph:
                el = etree.fromstring(MORPH % (NSDECL, NSDECL))
                self.stats['morph'] += 1
            else:
                el = etree.fromstring(FADE % NSDECL)
            sld.insert(idx, el)
            if s.anim:
                kinds = {}
                for el2 in s.tree.iter():
                    tag = etree.QName(el2).localname
                    if tag in ('sp', 'grpSp', 'pic', 'graphicFrame', 'cxnSp'):
                        c = el2.find('.//{%s}cNvPr' % NS['p'])
                        if c is not None:
                            kinds[int(c.get('id'))] = tag
                tb = TimingBuilder()
                x = tb.build(s.anim, kinds)
                if x:
                    sld.insert(idx + 1, etree.fromstring(x))
                    self.stats['anim_slides'] += 1
                    self.stats['effects'] += sum(len(st) for st in s.anim)
            if s.notes:
                s.ps.notes_slide.notes_text_frame.text = s.notes

    def anim_log(self):
        """[(n_diap, orden_clic, objeto, efecto, duración_ms, disparador)] para el guion de animaciones"""
        EFN = {'appear': 'Entrada · Aparecer', 'fade': 'Entrada · Desvanecer', 'wipe': 'Entrada · Barrido',
               'zoom': 'Entrada · Zoom', 'grow': 'Entrada · Zoom', 'float': 'Entrada · Flotar',
               'path': 'Trayectoria de movimiento', 'out': 'Salida · Desvanecer', 'outnow': 'Salida · Desaparecer',
               'pulse': 'Énfasis · Pulso', 'spin': 'Énfasis · Girar'}
        rows = []
        for n, s in enumerate(self.slides, 1):
            names = {int(c.get('id')): c.get('name') for c in s.tree.iter('{%s}cNvPr' % NS['p'])}
            for k, st in enumerate(s.anim, 1):
                for i, e in enumerate(st):
                    trig = 'Al hacer clic' if i == 0 else ('Después de la anterior' if e['after'] else 'Con la anterior')
                    ef = EFN[e['kind']]
                    if e['kind'] == 'wipe':
                        ef += ' (desde %s)' % {'l': 'izquierda', 'r': 'derecha', 't': 'arriba', 'b': 'abajo'}[e['dir']]
                    if e['kind'] == 'path':
                        ef += ' (Δx %.2f″, Δy %.2f″)' % (e['dx'], e['dy'])
                    rows.append((n, k, names.get(e['sid'], '?'), ef, e['dur'], trig))
        return rows

    def embed_fonts(self, fn):
        """incrusta Mali (OFL, incrustación instalable) como ppt/fonts/*.fntdata"""
        import zipfile, shutil
        tmp = fn + '.tmp'
        zin = zipfile.ZipFile(fn)
        zout = zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED)
        fonts = [('font1.fntdata', 'Mali-Regular.ttf', 'regular'), ('font2.fntdata', 'Mali-Bold.ttf', 'bold')]
        for it in zin.infolist():
            data = zin.read(it.filename)
            if it.filename == '[Content_Types].xml':
                data = data.replace(b'</Types>', b'<Default Extension="fntdata" ContentType="application/x-fontdata"/></Types>')
            elif it.filename == 'ppt/_rels/presentation.xml.rels':
                rels = ''.join('<Relationship Id="rIdF%d" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/font" Target="fonts/%s"/>' % (i + 1, f[0]) for i, f in enumerate(fonts))
                data = data.replace(b'</Relationships>', rels.encode() + b'</Relationships>')
            elif it.filename.startswith('ppt/theme/theme'):
                x = data.decode('utf8')
                i = x.find('<a:fontScheme')
                j = x.find('</a:fontScheme>')
                if i > 0:
                    x = x[:i] + x[i:j].replace('<a:latin typeface="Calibri"/>', '<a:latin typeface="Mali"/>') + x[j:]
                data = x.encode('utf8')
            elif it.filename == 'ppt/presentation.xml':
                x = data.decode('utf8')
                x = re.sub(r' saveSubsetFonts="\d"', '', x)
                x = x.replace(' type="screen4x3"', '')
                x = x.replace('<p:presentation ', '<p:presentation embedTrueTypeFonts="1" saveSubsetFonts="0" ', 1)
                import fuentes
                pan = fuentes.panose_hex(os.path.join(HERE, 'fuentes', 'Mali-Regular.ttf'))
                emb = ('<p:embeddedFontLst><p:embeddedFont><p:font typeface="Mali" panose="%s" pitchFamily="2" charset="0"/>%s</p:embeddedFont></p:embeddedFontLst>'
                       % (pan, ''.join('<p:%s r:id="rIdF%d"/>' % (f[2], i + 1) for i, f in enumerate(fonts))))
                i = x.index('<p:notesSz')
                j = x.index('/>', i) + 2
                x = x[:j] + emb + x[j:]
                data = x.encode('utf8')
            zout.writestr(it, data)
        import fuentes
        for f in fonts:
            # PowerPoint guarda y lee las fuentes incrustadas en formato EOT (Embedded OpenType)
            zout.writestr('ppt/fonts/' + f[0], fuentes.to_eot(os.path.join(HERE, 'fuentes', f[1])))
        zout.close(); zin.close()
        shutil.move(tmp, fn)

    def save(self, fn):
        self.finalize()
        self.prs.save(fn)
        self.embed_fonts(fn)
